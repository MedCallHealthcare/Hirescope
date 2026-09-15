from openpyxl import load_workbook
from datetime import datetime
from zipfile import BadZipFile
import json
import os
import shutil
import tempfile
import time


# PRODUCTION DAILY SCRAPING FILE
FILE_PATH = (
    r"C:\Users\HireScope\OneDrive - Med-Call Healthcare"
    r"\PH BPO Operations - Scraping Leads Tracker"
    r"\Daily Scraping.xlsx"
)


MAX_RETRY_ATTEMPTS = 5
RETRY_DELAY_SECONDS = 3


# Only these columns are required for the exporter
# to recognize a valid scraping worksheet.
#
# Other columns such as description, vertical,
# facility_type, score, ZoomInfo fields, etc.
# are supported automatically when they exist.
REQUIRED_HEADERS = {
    "job_title",
    "company_name",
    "location",
    "job_url",
}


def normalize_text(value):
    """
    Normalize text for duplicate comparisons:
    - convert None to empty string
    - lowercase
    - trim spaces
    - collapse multiple spaces
    """
    if value is None:
        return ""

    return " ".join(
        str(value).lower().strip().split()
    )


def normalize_header(value):
    """
    Normalize Excel column headers so the exporter
    can match columns even if capitalization,
    spaces, or hyphens are different.

    Examples:
    Job Title -> job_title
    Company Name -> company_name
    job_title -> job_title
    """
    if value is None:
        return ""

    return (
        str(value)
        .strip()
        .lower()
        .replace(" ", "_")
        .replace("-", "_")
    )


def get_header_map(worksheet):
    """
    Create a dictionary that maps each Excel header
    to its actual column number.

    Example:
    {
        "job_title": 1,
        "company_name": 2,
        "description": 6,
        "scraped_at": 7,
        "source": 8,
    }
    """
    header_map = {}

    for column in range(
        1,
        worksheet.max_column + 1
    ):
        header = worksheet.cell(
            row=1,
            column=column
        ).value

        normalized_header = normalize_header(
            header
        )

        if normalized_header:
            header_map[normalized_header] = column

    return header_map


def normalize_url(url):
    """Normalize a job URL for duplicate checking."""
    if not url:
        return ""

    return str(url).strip().lower().rstrip("/")


def create_unique_key(job):
    """
    Use job_url as the primary unique identifier.

    If a job has no URL, use:
    job_title + company_name + location
    """
    job_url = normalize_url(job.get("job_url"))

    if job_url:
        return f"url::{job_url}"

    title = normalize_text(job.get("job_title"))
    company = normalize_text(job.get("company_name"))
    location = normalize_text(job.get("location"))

    return f"fallback::{title}|{company}|{location}"


def find_target_worksheet(workbook):
    """
    Find a worksheet containing the required
    scraping headers.

    Extra spreadsheet columns are allowed and
    the required columns do not need to be in
    fixed positions.
    """
    for worksheet in workbook.worksheets:

        header_map = get_header_map(
            worksheet
        )

        available_headers = set(
            header_map.keys()
        )

        if REQUIRED_HEADERS.issubset(
            available_headers
        ):
            return worksheet, header_map

    return None, {}


def get_existing_keys(
    worksheet,
    header_map
):
    """
    Read all existing records and build duplicate
    keys using the actual Excel column positions.

    Completely empty rows are ignored.
    """
    existing_keys = set()
    existing_record_rows = []

    for row in range(
        2,
        worksheet.max_row + 1
    ):

        job = {
            "job_title": worksheet.cell(
                row=row,
                column=header_map["job_title"]
            ).value,

            "company_name": worksheet.cell(
                row=row,
                column=header_map["company_name"]
            ).value,

            "location": worksheet.cell(
                row=row,
                column=header_map["location"]
            ).value,

            "job_url": worksheet.cell(
                row=row,
                column=header_map["job_url"]
            ).value,
        }

        # Normalize the values before deciding
        # whether the row actually contains job data.
        has_data = any(
            normalize_text(value)
            for value in job.values()
        )

        if not has_data:
            continue

        key = create_unique_key(job)

        existing_keys.add(key)
        existing_record_rows.append(row)

    # Temporary diagnostic information.
    if existing_record_rows:
        print(
            "[DEBUG] Existing job data found in Excel rows: "
            f"{existing_record_rows[:20]}"
        )

        if len(existing_record_rows) > 20:
            print(
                "[DEBUG] Additional existing rows not shown: "
                f"{len(existing_record_rows) - 20}"
            )

    return existing_keys


def get_next_data_row(
    worksheet,
    header_map
):
    """
    Find the next real row for job data.

    worksheet.max_row cannot be trusted by itself because
    Excel may retain formatting or previously-used rows even
    after their contents have been deleted.

    If no existing job records are found, start at row 2.
    """

    required_columns = [
        header_map["job_title"],
        header_map["company_name"],
        header_map["location"],
        header_map["job_url"],
    ]

    for row in range(
        worksheet.max_row,
        1,
        -1
    ):

        row_has_data = False

        for column in required_columns:

            value = worksheet.cell(
                row=row,
                column=column
            ).value

            if normalize_text(value):
                row_has_data = True
                break

        if row_has_data:
            return row + 1

    return 2


def export_jobs_to_sharepoint(
    jobs,
    file_path=FILE_PATH,
    max_attempts=MAX_RETRY_ATTEMPTS,
    retry_delay=RETRY_DELAY_SECONDS,
):
    """
    Export processed jobs to a SharePoint-synced Excel workbook.

    Duplicate priority:
    1. job_url
    2. job_title + company_name + location

    Before saving, a temporary backup of the production workbook
    is created outside the OneDrive folder.

    The workbook is then saved directly to the existing OneDrive
    file instead of replacing it with a different file.

    After saving, the workbook is reopened and the newly-written
    rows are verified before the export is reported as successful.

    Temporary OneDrive/file-lock errors are automatically retried.
    """

    results = {
        "added": 0,
        "duplicates": 0,
        "errors": 0,
        "attempts": 0,
        "status": "pending",
        "message": "",
        "local_save_confirmed": False,
    }

    print("========================================")
    print("SHAREPOINT EXCEL EXPORT")
    print("========================================")

    if not jobs:
        print("No jobs provided.")

        results["status"] = "no_data"
        results["message"] = "No jobs were provided for export."

        return results

    print(f"Jobs received: {len(jobs)}")
    print(f"Excel file: {file_path}")
    print()

    last_error = None

    for attempt in range(1, max_attempts + 1):

        results["attempts"] = attempt

        workbook = None
        verification_workbook = None
        backup_file_path = None
        save_verified = False

        try:
            print(
                f"[EXPORT] Attempt "
                f"{attempt}/{max_attempts}"
            )

            # OneDrive may temporarily report that the file
            # does not exist while syncing.
            if not os.path.exists(file_path):
                raise FileNotFoundError(
                    f"Excel file not currently available: {file_path}"
                )

            # Always reopen the workbook on each retry so we
            # work with the newest synced version.
            workbook = load_workbook(file_path)

            worksheet, header_map = find_target_worksheet(
                workbook
            )

            if worksheet is None:
                workbook.close()
                workbook = None

                print(
                    "ERROR: Could not find a worksheet "
                    "with the required headers."
                )

                results["errors"] = len(jobs)
                results["status"] = "invalid_worksheet"
                results["message"] = (
                    "Could not find a worksheet containing "
                    "the required scraping column headers."
                )

                return results

            print(f"Worksheet: {worksheet.title}")

            print(
                f"Detected columns: "
                f"{len(header_map)}"
            )

            print(
                f"Worksheet max row reported by Excel: "
                f"{worksheet.max_row}"
            )

            existing_keys = get_existing_keys(
                worksheet,
                header_map
            )

            print(
                f"Existing records: "
                f"{len(existing_keys)}"
            )

            next_row_number = get_next_data_row(
                worksheet,
                header_map
            )

            print(
                f"Next actual write row: "
                f"{next_row_number}"
            )

            # Work with temporary counters first.
            # We only commit them to results after the
            # workbook has been successfully saved.
            added_count = 0
            duplicate_count = 0

            # Track the rows written during this export
            # so they can be verified after saving.
            written_records = []

            for job in jobs:

                unique_key = create_unique_key(job)

                if unique_key in existing_keys:

                    duplicate_count += 1

                    print(
                        f"DUPLICATE: "
                        f"{job.get('job_title', '')} | "
                        f"{job.get('company_name', '')}"
                    )

                    continue

                scraped_at = job.get("scraped_at")

                if not scraped_at:
                    scraped_at = datetime.now().strftime(
                        "%Y-%m-%d %H:%M:%S"
                    )

                # Copy the processed job dictionary so
                # the original scraped record is not changed.
                job_data = dict(job)

                job_data["scraped_at"] = scraped_at

                # Use the next available Excel row.
                row_number = next_row_number

                # Match job fields to the spreadsheet
                # columns using their header names.
                for (
                    header_name,
                    column_number
                ) in header_map.items():

                    if header_name not in job_data:
                        continue

                    value = job_data.get(
                        header_name,
                        ""
                    )

                    # Excel cells cannot directly store
                    # lists, tuples, or sets.
                    if isinstance(
                        value,
                        (list, tuple, set)
                    ):
                        value = ", ".join(
                            str(item)
                            for item in value
                        )

                    # Excel cells also cannot directly
                    # store Python dictionaries.
                    elif isinstance(value, dict):
                        value = json.dumps(
                            value,
                            ensure_ascii=False
                        )

                    worksheet.cell(
                        row=row_number,
                        column=column_number,
                        value=value
                    )

                # Prevent duplicates within the same batch.
                existing_keys.add(unique_key)

                added_count += 1

                # Remember exactly where this record was written
                # so the saved workbook can be verified later.
                written_records.append(
                    (
                        row_number,
                        unique_key,
                    )
                )

                # IMPORTANT:
                # Move to a new Excel row after EACH
                # successfully-added job.
                next_row_number += 1

                print(
                    f"ADDED: "
                    f"{job.get('job_title', '')} | "
                    f"{job.get('company_name', '')} "
                    f"-> Excel row {row_number}"
                )

            # Nothing new needs to be written.
            if added_count == 0:

                workbook.close()
                workbook = None

                results["duplicates"] = duplicate_count
                results["status"] = "no_changes"
                results["message"] = (
                    "All jobs already exist in the workbook."
                )
                results["local_save_confirmed"] = True

                print()
                print("No new jobs to add.")
                print(
                    f"Duplicates: "
                    f"{duplicate_count}"
                )

                return results

            # Create a temporary backup OUTSIDE the OneDrive
            # folder before modifying the production workbook.
            #
            # This protects the existing production data without
            # creating extra backup files inside SharePoint.
            timestamp = datetime.now().strftime(
                "%Y%m%d_%H%M%S_%f"
            )

            backup_file_path = os.path.join(
                tempfile.gettempdir(),
                f"hirescope_daily_scraping_backup_{timestamp}.xlsx"
            )

            shutil.copy2(
                file_path,
                backup_file_path
            )

            print(
                f"[BACKUP] Temporary backup created: "
                f"{backup_file_path}"
            )

            # Save DIRECTLY to the existing OneDrive workbook.
            #
            # Do not replace the OneDrive file with another file.
            workbook.save(file_path)
            workbook.close()
            workbook = None

            print(
                "[SAVE] Workbook saved directly to "
                "the OneDrive file."
            )

            # Give Windows / OneDrive a moment to process
            # the local file modification.
            time.sleep(2)

            # Reopen the exact production workbook we just saved.
            verification_workbook = load_workbook(
                file_path
            )

            (
                verification_worksheet,
                verification_header_map,
            ) = find_target_worksheet(
                verification_workbook
            )

            if verification_worksheet is None:
                raise RuntimeError(
                    "Save verification failed: "
                    "the scraping worksheet could not be found "
                    "after reopening the workbook."
                )

            # Verify every newly-written row individually.
            for (
                written_row,
                expected_key,
            ) in written_records:

                verification_job = {
                    "job_title": verification_worksheet.cell(
                        row=written_row,
                        column=verification_header_map["job_title"]
                    ).value,

                    "company_name": verification_worksheet.cell(
                        row=written_row,
                        column=verification_header_map["company_name"]
                    ).value,

                    "location": verification_worksheet.cell(
                        row=written_row,
                        column=verification_header_map["location"]
                    ).value,

                    "job_url": verification_worksheet.cell(
                        row=written_row,
                        column=verification_header_map["job_url"]
                    ).value,
                }

                actual_key = create_unique_key(
                    verification_job
                )

                if actual_key != expected_key:
                    raise RuntimeError(
                        "Save verification failed at "
                        f"Excel row {written_row}."
                    )

                print(
                    f"[VERIFY] Excel row "
                    f"{written_row}: OK"
                )

            verification_workbook.close()
            verification_workbook = None

            # Mark the save as verified only after every
            # newly-written row has been found successfully.
            save_verified = True

            results["added"] = added_count
            results["duplicates"] = duplicate_count
            results["errors"] = 0
            results["status"] = "success"
            results["message"] = (
                "Jobs were successfully written and verified "
                "in the local OneDrive workbook."
            )
            results["local_save_confirmed"] = True

            print()
            print("========================================")
            print("EXPORT COMPLETE")
            print("========================================")
            print(
                f"Added:      "
                f"{results['added']}"
            )
            print(
                f"Duplicates: "
                f"{results['duplicates']}"
            )
            print(
                f"Errors:     "
                f"{results['errors']}"
            )
            print(
                f"Attempts:   "
                f"{results['attempts']}"
            )
            print()
            print(
                "Local workbook save confirmed."
            )
            print(
                "OneDrive will sync the saved file "
                "to SharePoint asynchronously."
            )

            return results

        except (
            PermissionError,
            FileNotFoundError,
            BadZipFile,
            OSError,
        ) as error:

            last_error = error

            print()
            print(
                f"[WARN] Export attempt "
                f"{attempt} failed:"
            )
            print(f"       {error}")

        except Exception as error:

            last_error = error

            print()
            print(
                f"[ERROR] Unexpected export error: "
                f"{error}"
            )

        finally:

            if workbook is not None:
                try:
                    workbook.close()
                except Exception:
                    pass

            if verification_workbook is not None:
                try:
                    verification_workbook.close()
                except Exception:
                    pass

            # If a backup was created, either:
            #
            # 1. delete it after a verified successful save, or
            # 2. restore it if the save/verification failed.
            if (
                backup_file_path
                and os.path.exists(backup_file_path)
            ):

                if save_verified:

                    try:
                        os.remove(
                            backup_file_path
                        )

                        print(
                            "[BACKUP] Temporary backup removed "
                            "after successful verification."
                        )

                    except OSError:
                        pass

                else:

                    restore_succeeded = False

                    try:
                        shutil.copy2(
                            backup_file_path,
                            file_path
                        )

                        restore_succeeded = True

                        print(
                            "[RESTORE] Original workbook restored "
                            "from backup."
                        )

                    except Exception as restore_error:

                        print(
                            "[RESTORE ERROR] Could not restore "
                            "the original workbook:"
                        )

                        print(
                            f"                "
                            f"{restore_error}"
                        )

                        print(
                            "[IMPORTANT] Backup retained at:"
                        )

                        print(
                            f"            "
                            f"{backup_file_path}"
                        )

                    if restore_succeeded:
                        try:
                            os.remove(
                                backup_file_path
                            )
                        except OSError:
                            pass

        if attempt < max_attempts:

            print(
                f"[INFO] Waiting "
                f"{retry_delay} seconds before retry..."
            )

            time.sleep(retry_delay)

    # All attempts failed.
    results["errors"] = len(jobs)
    results["status"] = "failed"
    results["message"] = (
        f"Export failed after "
        f"{max_attempts} attempts. "
        f"Last error: {last_error}"
    )

    print()
    print("========================================")
    print("EXPORT FAILED")
    print("========================================")
    print(
        f"Attempts: {results['attempts']}"
    )
    print(
        f"Error: {last_error}"
    )

    return results